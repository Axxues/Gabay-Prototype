"""Check THESIS-GABAY-v5.docx Table 3/6 benchmark totals equal 140 (Task 5)."""
import sys
from pathlib import Path

DOCX = Path(__file__).resolve().parent.parent / "THESIS-GABAY-v5.docx"

EXPECTED_REASONING = ["Reasoning Capability", "4", "7", "6", "3", "20"]
EXPECTED_TOTAL = ["Total", "32", "45", "36", "27", "140"]

errors: list[str] = []


def norm(values: list[str]) -> list[str]:
    return [" ".join(v.split()) for v in values]


def main() -> int:
    try:
        import docx
    except ImportError:
        print("FAIL: python-docx not installed")
        return 1
    if not DOCX.exists():
        print(f"FAIL: missing {DOCX}")
        return 1
    doc = docx.Document(str(DOCX))

    target = None
    for t in doc.tables:
        if len(t.rows) == 0 or len(t.rows[0].cells) != 6:
            continue
        header = norm([c.text for c in t.rows[0].cells])
        if header[0] == "Performance Sub-Objective" and header[5] == "Total Queries":
            target = t
            break
    if target is None:
        errors.append("benchmark distribution table (Performance Sub-Objective/Total Queries) not found")
    else:
        rows = {norm([c.text for c in r.cells])[0]: norm([c.text for c in r.cells]) for r in target.rows}
        for label, expected in (("Reasoning Capability", EXPECTED_REASONING), ("Total", EXPECTED_TOTAL)):
            actual = rows.get(label)
            if actual != expected:
                errors.append(f"{label} row: expected {expected}, got {actual}")

    for i, p in enumerate(doc.paragraphs):
        if "155" in p.text:
            errors.append(f"paragraph P{i} still contains '155': {p.text[:120]!r}")
        if "fifty-five" in p.text.lower():
            errors.append(f"paragraph P{i} still contains 'fifty-five': {p.text[:120]!r}")

    if errors:
        print("FAIL: thesis check failed")
        for e in errors:
            print(f" - {e}")
        return 1
    print("PASS: thesis Table(140): reasoning 4/7/6/3/20, total 32/45/36/27/140; no 155/fifty-five left")
    return 0


if __name__ == "__main__":
    sys.exit(main())
